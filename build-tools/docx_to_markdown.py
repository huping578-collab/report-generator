#!/usr/bin/env python3
"""Convert Word report templates under templates/ to Markdown skeleton files.

The Markdown forms are codepage-safe, git-friendly and ship inside the
installer; image media are referenced with placeholders (not copied), so the
resulting .md is small (the 111MB Chongqing template collapses to ~tens of KB).
Usage: python build-tools/docx_to_markdown.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"
TARGETS = ["重庆项目报告模板.docx", "广东项目第五章模板.docx"]
HEADING_MARKERS = {1: "#", 2: "##", 3: "###", 4: "####", 5: "#####", 6: "######"}


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style.name.lower() if paragraph.style is not None else ""
    for marker, level in (("heading", 1), ("标题", 1)):
        if style.startswith(marker):
            m = re.search(r"(\d+)$", style)
            if m:
                return int(m.group(1))
            return 1
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.outlineLvl is not None:
        return int(p_pr.outlineLvl.val) + 1
    return None


def _inline_text(paragraph: Paragraph) -> str:
    parts = []
    for run in paragraph.runs:
        if run._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"):
            parts.append(_picture_mark(paragraph))
            break
        parts.append(run.text or "")
    text = "".join(parts).strip()
    return text


def _picture_mark(paragraph: Paragraph) -> str:
    images = []
    for rel in paragraph.part.rels.values():
        if rel.reltype.endswith("/image") and rel.target_part.partname:
            images.append(str(rel.target_part.partname).split("/")[-1])
    if not images:
        return "![图片](media/图片.png)"
    return "\n".join(f"![{name}](media/{name})" for name in images)


def _render_table(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append(text.replace("|", "\\|"))
        if len(set(cells)) == 1 and all(c == "" for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return []
    width = max(len(r) for r in rows)

    def fmt(row: list[str]) -> str:
        cell = row + [""] * (width - len(row))
        return "| " + " | ".join(cell) + " |"

    lines = [fmt(rows[0]), "|" + "---|" * width]
    for r in rows[1:]:
        lines.append(fmt(r))
    return lines


def convert(path: Path) -> str:
    doc = Document(str(path))
    body = doc.element.body
    lines: list[str] = []
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            level = _heading_level(paragraph)
            if level is not None and text:
                lines.append(f"{HEADING_MARKERS.get(level, '#')} {text}")
                lines.append("")
            elif text:
                lines.append(text)
                lines.append("")
            elif "inline" in child.xml:
                pic = _picture_mark(paragraph)
                if pic:
                    lines.append(pic)
                    lines.append("")
        elif tag == "tbl":
            lines.extend(_render_table(Table(child, doc)))
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def main() -> int:
    for name in TARGETS:
        source = TEMPLATES / name
        if not source.is_file():
            print(f"skip: {name} not found", flush=True)
            continue
        target = source.with_suffix(".md")
        markdown = convert(source)
        target.write_text(markdown, encoding="utf-8")
        print(f"{name} -> {target.name} ({len(markdown)} chars)", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
