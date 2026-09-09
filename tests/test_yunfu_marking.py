"""云浮宽表输入、共享区段和报告/Excel持久回归（不复制参考统计数）。"""
from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import hashlib
from docx import Document

import openpyxl
from backend import report_engine as engine

HEADERS = ["地市", "路线编号", "检测方向", "管养单位", "桩号",
           "主车道左侧标线逆反亮度系数", "主车道右侧标线逆反亮度系数",
           "左侧标线逆反射目标值", "右侧标线逆反射目标值", "计算区间"]
TEMPLATE = Path(__file__).resolve().parents[1] / "templates/广东项目第五章模板.md"
REAL_INPUT = Path("C:/文件/工作工具台/广东报告数据/标线数据/云浮市标线统计.xlsx")


def wide_row(station="1000+000.0-1000+020.0", **changes):
    row = dict(zip(HEADERS, ["云浮", "G234", "上行", "甲管养单位", station, 40, 90, 50, 80, 20]))
    row.update(changes)
    return row


def scan_rows(rows, headers=HEADERS):
    with TemporaryDirectory() as tmp:
        path = Path(tmp) / "宽表.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append([row.get(key) for key in headers])
        wb.save(path)
        wb.close()
        return engine.GuangdongInputScanner(tmp)._process_file(path)


class YunfuMarkingTest(unittest.TestCase):
    def test_wide_input_preserves_each_side_and_source_coordinates(self):
        rows = [wide_row(), wide_row("1000+020.0-1000+040.0", **{HEADERS[5]: None}),
                wide_row("1000+040.0-1000+060.0", **{HEADERS[6]: None, "检测方向": "下行"})]
        scanned = scan_rows(rows)
        self.assertEqual(scanned["issues"], [])
        got = [(r.get("source_row"), r.get("source_range"), r["station_m"], r.get("end_m"),
                r["direction"], r["marking_position"], r["value"], r["target"], r["manager"])
               for r in scanned["marking"]]
        expected = []
        for row_no, row in enumerate(rows, 2):
            a, b = row["桩号"].split("-")
            start, end = [float(x.split("+")[0]) * 1000 + float(x.split("+")[1]) for x in (a, b)]
            for vi, ti, side in ((5, 7, "左侧标线"), (6, 8, "右侧标线")):
                if row[HEADERS[vi]] is not None:
                    expected.append((row_no, row["桩号"], start, end, row["检测方向"], side,
                                     row[HEADERS[vi]], row[HEADERS[ti]], row["管养单位"]))
        self.assertEqual(got, expected)

    def test_invalid_metadata_and_numbers_are_reported_without_fabrication(self):
        for field in ("地市", "路线编号", "检测方向", "管养单位", "桩号"):
            with self.subTest(field=field):
                result = scan_rows([wide_row(**{field: " "})])
                self.assertFalse(result["marking"])
                self.assertTrue(result["issues"])
                self.assertIn("第2行", "".join(result["issues"]))
        for column in (HEADERS[5], HEADERS[7]):
            for value in ("NaN", "inf", "-inf", -1, "not-a-number"):
                with self.subTest(column=column, value=value):
                    result = scan_rows([wide_row(**{column: value})])
                    self.assertEqual([r["marking_position"] for r in result["marking"]], ["右侧标线"])
                    self.assertTrue(result["issues"])
        result = scan_rows([wide_row(**{HEADERS[7]: None})])
        self.assertEqual([r["target"] for r in result["marking"]], [80])
        self.assertTrue(result["issues"])
        for value in (float("nan"), float("inf"), -float("inf")):
            self.assertIsNone(engine._float(value))

    def test_city_labels_differing_only_by_shi_suffix_merge_into_one_bundle(self):
        rows = [wide_row(), wide_row(**{"地市": "云浮市", "路线编号": "G2518"})]
        scanned = scan_rows(rows)
        route_index = engine.RouteCategoryIndex({("云浮", "G234"): "普通国省道", ("云浮市", "G2518"): "高速公路"})
        bundles = engine.GuangdongBatchRunner.build_bundles(scanned, route_index)
        self.assertEqual(list(bundles), ["云浮市"])
        self.assertEqual(sorted({r["route"] for r in bundles["云浮市"]["marking"]}), ["G234", "G2518"])
        self.assertTrue(all(r["city"] == "云浮市" for r in bundles["云浮市"]["marking"]))

    def test_legacy_single_value_keeps_segment_and_default_target(self):
        scanner = engine.GuangdongInputScanner(".", engine.RouteCategoryIndex({("云浮", "G234"): "普通国省道"}))
        row = {"计算区间": "K1+000～K1+020", "逆反亮度系数": 88}
        record = scanner._convert("marking", row, Path("G234下行标线3K1K2.csv"), "CSV")
        self.assertEqual((record["city"], record["direction"], record["marking_position"], record["target"]),
                         ("云浮", "下行", "标线3", 80))
        self.assertEqual(record["segment"], "G234下行K1～K2")
        self.assertEqual(record["source_range"], row["计算区间"])
        bundle = engine.GuangdongBatchRunner.build_bundles({"marking": [record]}, scanner.route_index)["云浮"]
        self.assertEqual(bundle["marking"][0]["segment"], record["segment"])
        row["逆反亮度系数目标值"] = -1
        with self.assertRaises(ValueError):
            scanner._convert("marking", row, Path("G234下行标线3K1K2.csv"), "CSV")

    def test_only_connected_ranges_share_report_segment(self):
        rows = [wide_row(), wide_row("1000+020.0-1000+040.0", **{"计算区间": 10}),
                wide_row("1000+035.0-1000+060.0"), wide_row("1000+080.0-1000+100.0"),
                wide_row(**{"管养单位": "乙管养单位"}), wide_row(**{"路线编号": "S265"}),
                wide_row("1000+020.0-1000+000.0", **{"检测方向": "下行"}),
                wide_row("1000+040.0-1000+020.0", **{"检测方向": "下行"}),
                wide_row(**{"地市": "肇庆"})]
        scanned = scan_rows(rows)
        route_index = engine.RouteCategoryIndex({(r["地市"], r["路线编号"]): "普通国省道" for r in rows})
        bundles = engine.GuangdongBatchRunner.build_bundles(scanned, route_index)
        marking = bundles["云浮"]["marking"]
        groups = engine.GuangdongStatistics.marking_segment_pair_summary(marking)
        self.assertEqual(len(groups), 5)
        self.assertEqual(len(engine.GuangdongStatistics.marking_segment_summary(marking)), 10)
        first = [r for r in marking if r["source_row"] in (2, 3, 4)]
        self.assertEqual(len({r["segment"] for r in first}), 1)
        gap = next(r for r in marking if r["source_row"] == 5)
        self.assertNotEqual(first[0]["segment"], gap["segment"])
        self.assertEqual({r["source_range"] for r in first}, {r["桩号"] for r in rows[:3]})
        down = [r for r in marking if r["direction"] == "下行"]
        self.assertEqual(len({r["segment"] for r in down}), 1)
        self.assertEqual(down[0]["station_m"], 1000020)
        self.assertEqual(down[0]["end_m"], 1000000)
        self.assertEqual(len(marking), 16)
        all_rows = marking + bundles["肇庆"]["marking"]
        self.assertEqual(len(engine.GuangdongStatistics.marking_segment_pair_summary(all_rows)), 6)
        self.assertEqual(len(scanned["marking"]), len(all_rows))

    def test_explicit_left_right_names_never_swap(self):
        self.assertEqual(engine.marking_side_names(["右侧标线", "左侧标线"]),
                         {"左侧标线": "左侧标线", "右侧标线": "右侧标线"})
        self.assertEqual(engine.marking_side_names(["右侧标线"]), {"右侧标线": "右侧标线"})
        self.assertEqual(engine.marking_side_names(["标线2", "标线3"]),
                         {"标线2": "左侧标线", "标线3": "右侧标线"})
        self.assertEqual(engine.marking_side_names(["标线1", "标线2"]),
                         {"标线1": "左侧标线", "标线2": "右侧标线"})

    def test_word_charts_use_gd03_charts_and_excel_export_keeps_groups_isolated(self):
        rows = []
        for changes in ({}, {"管养单位": "乙管养单位"}, {"路线编号": "S265"}, {"检测方向": "下行"}):
            rows.extend([wide_row(**changes), wide_row("1000+020.0-1000+040.0", **changes)])
        rows.append(wide_row("1000+080.0-1000+100.0", **{HEADERS[5]: None}))
        scanned = scan_rows(rows)
        route_index = engine.RouteCategoryIndex({("云浮", r["路线编号"]): "普通国省道" for r in rows})
        bundle = engine.GuangdongBatchRunner.build_bundles(scanned, route_index)["云浮"]
        with TemporaryDirectory() as tmp:
            images = engine.guangdong_report_images(bundle, tmp)
            self.assertEqual(len(images), 9)  # 4 identities x 2 sides + right-only gap（Excel 逐区段图不变）
            doc = Document(engine.GuangdongChapterWriter.write("云浮", bundle, tmp, TEMPLATE, {"marking": 5, "height": 5, "bolt": 5}))
            chart_dir = Path(tmp) / "云浮" / "charts"
            generated = {hashlib.sha256(p.read_bytes()).hexdigest() for p in chart_dir.glob("*gd03_*.png")}
            embedded = {hashlib.sha256(doc.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest()
                        for shape in doc.inline_shapes}
            self.assertTrue(generated, "未生成 GD03 统计图")
            self.assertTrue(embedded and embedded <= generated, "docx 内嵌图必须来自本次生成的 GD03 统计图")
            table = next(t for t in doc.tables if "检测范围" in [c.text for c in t.rows[0].cells])
            headers = [c.text for c in table.rows[0].cells]
            self.assertIn("管养单位", headers)
            self.assertLess(headers.index("左侧合格率(%)"), headers.index("右侧合格率(%)"))
            self.assertEqual(len(table.rows), 4)  # 表头 + 3 个“路线—管养单位”100m 单元
            for cells in table.rows[1:]:
                record = dict(zip(headers, (c.text for c in cells.cells)))
                self.assertNotIn("None", [c.text for c in cells.cells])
                self.assertEqual(record["左侧合格率(%)"], "0.0")
                self.assertEqual(record["右侧合格率(%)"], "100.0")
                self.assertEqual(record["总体合格率(%)"], "50.0")
            wb = openpyxl.load_workbook(engine.write_guangdong_chart_workbook(bundle, tmp))
            ws = wb["标线逆反射"]
            self.assertEqual(len(ws._charts), 9)
            exported = []
            header = None
            for values in ws.iter_rows(values_only=True):
                if values[0] == "桩号":
                    header = list(values)
                elif values[0]:
                    exported.append(dict(zip(header, values)))
            def key(r):
                return (r["路线"], r["方向"], r["管养单位"], r["标线位置"], r["源行号"],
                        r["源桩号范围"], r["逆反射亮度系数"], r["目标值"])
            expected = [(r["route"], r["direction"], r["manager"], r["marking_position"], r["source_row"],
                         r["source_range"], r["value"], r["target"]) for r in bundle["marking"]]
            self.assertEqual(Counter(map(key, exported)), Counter(expected))
            # 每张图的实测/目标系列引用各自数据块，且没有互相覆盖。
            from openpyxl.utils.cell import range_boundaries
            count = 0
            for chart in ws._charts:
                self.assertEqual(len(chart.series), 2)
                ref = chart.series[0].val.numRef.f.split("!")[1]
                x1, y1, x2, y2 = range_boundaries(ref)
                self.assertEqual(x1, x2)
                points = [ws.cell(y, x1).value for y in range(y1, y2 + 1)]
                self.assertTrue(all(isinstance(v, (float, int)) for v in points))
                count += len(points)
            self.assertEqual(count, len(bundle["marking"]))
            wb.close()

    def test_continuous_marking_weak_uses_cell_end_and_length(self):
        def cell(offset):
            return {"route": "G999", "direction": "上行", "marking_position": "左侧标线",
                    "station_m": 1000.0 + 20 * offset, "end_m": 1020.0 + 20 * offset,
                    "value": 10.0, "target": 80.0}

        rows = [cell(i) for i in range(6)]  # 6 个 20m 单元，连续区间恰为 120m
        weak = engine.GuangdongStatistics.continuous_marking_weak(rows, minimum_length_m=120)
        self.assertEqual(weak, [{"route": "G999", "direction": "上行", "marking_position": "左侧标线",
                                 "start_m": 1000.0, "end_m": 1120.0}])
        # 断点拆段：两段各 40m，均不足阈值
        broken = [cell(0), cell(1), dict(cell(2), value=90.0), cell(3), cell(4)]
        self.assertEqual(engine.GuangdongStatistics.continuous_marking_weak(broken, minimum_length_m=120), [])
        # 下行逆向桩号：起点取区间低端、止点取区间高端
        reverse = [{"route": "G998", "direction": "下行", "marking_position": "右侧标线",
                    "station_m": 2000.0 - 20 * i, "end_m": 2000.0 - 20 * i - 20.0,
                    "value": 10.0, "target": 80.0} for i in range(6)]
        weak = engine.GuangdongStatistics.continuous_marking_weak(reverse, minimum_length_m=120)
        self.assertEqual([(w["start_m"], w["end_m"]) for w in weak], [(1880.0, 2000.0)])

    def test_bolt_over5_runs_uses_per_km_keys(self):
        def bolt(km, splice_missing, connection_missing):
            return {"route": "G999", "direction": "上行", "manager": "甲管养单位",
                    "station_m": km * 1000.0, "splice": 100.0, "splice_missing": float(splice_missing),
                    "connection": 100.0, "connection_missing": float(connection_missing), "outline": 0.0}

        rows = [bolt(1000, 10, 5), bolt(1001, 8, 4), bolt(1002, 0, 0)]
        runs = engine.GuangdongStatistics.bolt_over5_runs(rows)
        self.assertEqual(len(runs), 1)
        run = runs[0]
        self.assertEqual((run["start_m"], run["end_m"], run["length_km"]), (1000000.0, 1002000.0, 2))
        self.assertEqual(run["missing"], 27.0)
        self.assertAlmostEqual(run["rate"], 27.0 / (215.0 + 212.0), places=12)

    def test_height_over10_runs_buckets_per_kind(self):
        """R-02：同公里内两波与三波分别判定，不得因混合稀释而漏报（S51 K37 口径）。"""
        rows = [{"route": "S51", "direction": "下行", "manager": "甲管养单位", "station_m": 37000.0 + index,
                 "guardrail_type": "二波", "height": 600.0} for index in range(131)]
        rows += [{"route": "S51", "direction": "下行", "manager": "甲管养单位", "station_m": 37000.0 + index,
                  "guardrail_type": "三波", "height": 850.0 if index < 5 else 700.0} for index in range(20)]
        runs = engine.GuangdongStatistics.height_over10_runs(rows)
        self.assertEqual(len(runs), 1)
        run = runs[0]
        self.assertEqual((run["route"], run["direction"], run["kind"]), ("S51", "下行", "三波"))
        self.assertEqual((run["count"], run["over_count"]), (20, 5))
        self.assertAlmostEqual(run["over_ratio"], 0.25)
        self.assertEqual((run["start_m"], run["end_m"]), (37000.0, 38000.0))


if __name__ == "__main__":
    unittest.main()
