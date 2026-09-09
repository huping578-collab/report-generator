"""广东第五章结构回归；最小输入仅验证结构，不使用参考报告的实测数据。"""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from docx.oxml.ns import qn
from backend.report_engine import GuangdongChapterWriter


class YunfuTemplateTest(unittest.TestCase):
    def test_generated_chapter_structure(self):
        template = Path(__file__).resolve().parents[1] / "templates/广东项目第五章模板.md"
        bundle = {
            "comparison_detail": [
                {"indicator": "marking", "city": "结构测试", "route": "G2518", "direction": "上行",
                 "segment": "K1+000~K2+000", "category": "高速公路", "source": "正式检测后",
                 "manual": 100.0, "automatic": 101.0},
            ],
            "weak_segments": [
                {"route": "G2518", "direction": "下行", "segment": "G2518下行K0343～K0308",
                 "type": "护栏高度偏差超10cm", "reason": "三波护栏偏差超10 cm点数1个"},
            ],
        }
        with TemporaryDirectory() as output:
            path = GuangdongChapterWriter.write(
                "结构测试", bundle, output, template,
                {"marking": 5, "height": 5, "bolt": 5},
            )
            doc = Document(path)
        headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
        sections = [text for style, text in headings if style == "Heading 2"]
        self.assertEqual(sections, ["（一）高速公路交安设施技术状况", "（二）普通国省道交安设施技术状况", "（三）工作建议"])
        texts = [text for _, text in headings]
        for title in ("1.沿线设施技术状况TCI", "（5）防眩设施状况", "2.标线、护栏自动化检测", "（1）标线逆反射亮度系数情况", "（2）波形梁护栏中心高度情况", "（3）螺栓缺失情况"):
            self.assertEqual(texts.count(title), 2, title)
        expected = ["1.重点路段处治建议（如有）", "2.迎国评工作建议", "3.养护提升建议"]
        self.assertEqual([text for style, text in headings if style == "Heading 3"][-3:], expected)
        for section in sections[:2]:
            start = texts.index(section)
            self.assertLess(texts.index("（5）防眩设施状况", start), texts.index("2.标线、护栏自动化检测", start))
        captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
        self.assertTrue(captions)
        self.assertTrue(all(text.startswith(("表5-", "图5-")) for text in captions))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("本次未提供人工复核对比记录", text)
        self.assertIn("本次未识别到高速公路优先处治路段", text)
        self.assertNotIn("绿化管护状况", text)
        self.assertNotIn("细化展开", text)
        body_sizes = {run.font.size.pt for p in doc.paragraphs if p.style.name == "Normal" and p.text.strip()
                      for run in p.runs if run.font.size}
        table_sizes = {run.font.size.pt for table in doc.tables for row in table.rows for cell in row.cells
                       for p in cell.paragraphs for run in p.runs if run.font.size}
        caption_sizes = {run.font.size.pt for p in doc.paragraphs if p.style.name == "Caption"
                         for run in p.runs if run.font.size}
        self.assertEqual(body_sizes, {12})
        self.assertEqual(table_sizes, {9})
        self.assertEqual(caption_sizes, {9})
        header_tr_pr = doc.tables[0].rows[0]._tr.find(qn('w:trPr'))
        self.assertIsNotNone(header_tr_pr)
        self.assertIsNotNone(header_tr_pr.find(qn('w:tblHeader')))


if __name__ == "__main__":
    unittest.main()
